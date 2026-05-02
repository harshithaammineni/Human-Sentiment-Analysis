from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "EEEM068_Human_Sentiment_Analysis_Report_clean_rewrite.docx"
FIG_OUT = ROOT / "report" / "clean_generated_figures"
OUTPUTS = ROOT / "outputs"
METRICS = OUTPUTS / "metrics"
PLOTS = OUTPUTS / "plots"
CMS = OUTPUTS / "confusion_matrices"
DATA = ROOT / "data" / "processed"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def font(size=24, bold=False):
    name = "timesbd.ttf" if bold else "times.ttf"
    candidates = [
        Path("C:/Windows/Fonts") / name,
        Path("C:/Windows/Fonts/times.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def wrap(draw, text, fnt, max_w):
    words = str(text).split()
    lines, cur = [], ""
    for word in words:
        trial = (cur + " " + word).strip()
        if draw.textlength(trial, font=fnt) <= max_w or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines or [""]


def table_png(name: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> Path:
    FIG_OUT.mkdir(parents=True, exist_ok=True)
    path = FIG_OUT / f"{name}.png"
    W = 1500
    margin = 24
    usable = W - 2 * margin
    sw = sum(widths)
    col_w = [int(usable * w / sw) for w in widths]
    col_w[-1] += usable - sum(col_w)
    fh, fb, fbb = font(30, True), font(28), font(28, True)
    pad = 12

    dummy = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    heights = []
    for row, fnt in [(headers, fh)] + [(r, fb) for r in rows]:
        max_lines = max(len(wrap(dummy, txt, fnt, cw - 2 * pad)) for txt, cw in zip(row, col_w))
        heights.append(max(54, max_lines * (fnt.size + 5) + 2 * pad))

    H = 2 * margin + sum(heights)
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    y = margin
    colors = {"navy": (26, 54, 93), "grid": (200, 210, 220), "light": (244, 247, 250), "ink": (22, 28, 36)}

    def draw_row(vals, h, fill, text_color, fnt, first_bold=False):
        nonlocal y
        x = margin
        for j, (txt, cw) in enumerate(zip(vals, col_w)):
            d.rectangle((x, y, x + cw, y + h), fill=fill, outline=colors["grid"], width=3)
            lines = wrap(d, txt, fnt, cw - 2 * pad)
            line_h = fnt.size + 5
            ty = y + (h - len(lines) * line_h) / 2
            use_font = fbb if first_bold and j == 0 else fnt
            for line in lines:
                if j == 0:
                    tx = x + pad
                else:
                    tx = x + (cw - d.textlength(line, font=use_font)) / 2
                d.text((tx, ty), line, font=use_font, fill=text_color)
                ty += line_h
            x += cw
        y += h

    draw_row(headers, heights[0], colors["navy"], "white", fh)
    for i, row in enumerate(rows):
        draw_row(row, heights[i + 1], colors["light"] if i % 2 else "white", colors["ink"], fb, True)
    img.save(path)
    return path


def comparison_bar(rows: list[tuple[str, float]]) -> Path:
    FIG_OUT.mkdir(parents=True, exist_ok=True)
    path = FIG_OUT / "macro_f1_compact.png"
    W, H = 1500, 850
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    navy = (26, 54, 93)
    teal = (38, 125, 142)
    coral = (181, 85, 60)
    d.text((40, 20), "Test macro-F1 by model", font=font(42, True), fill=navy)
    x0, y0, bw, bh, gap = 500, 95, 850, 48, 30
    for i, (label, val) in enumerate(rows):
        y = y0 + i * (bh + gap)
        d.text((40, y + 8), label, font=font(27, "CLIP" in label), fill=(20, 20, 20))
        color = coral if "CLIP" in label else teal
        w = int(bw * val / 0.60)
        d.rounded_rectangle((x0, y, x0 + w, y + bh), radius=10, fill=color)
        d.text((x0 + w + 14, y + 8), f"{val:.3f}", font=font(27, True), fill=(20, 20, 20))
    d.line((x0, y0 - 12, x0, y0 + len(rows) * (bh + gap) - gap + 12), fill=(210, 218, 226), width=3)
    img.save(path)
    return path


def set_columns(section, num=2, space=360):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space))


def style_doc(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(8.7)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def p(doc, text="", size=8.7, bold=False, italic=False, align=None, keep=False):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = 1.0
    if keep:
        para.paragraph_format.keep_with_next = True
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return para


def heading(doc, text, level=1):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4 if level == 1 else 2)
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.keep_with_next = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9.3 if level == 1 else 8.8)
    run.bold = True
    return para


def add_fig(doc, path: Path, caption: str, width=3.35):
    if not path.exists():
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(1)
    para.add_run().add_picture(str(path), width=Inches(width))
    cap = p(doc, caption, size=7.2, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.space_after = Pt(3)


def add_table_fig(doc, path: Path, caption: str, width=3.35):
    add_fig(doc, path, caption, width)


def metric_rows():
    face = read_json(METRICS / "face_resnet18_metrics.json")
    face_deg = read_json(METRICS / "face_resnet18_test_degraded_metrics.json")
    face_aug = read_json(METRICS / "face_resnet18_aug_test_original_metrics.json")
    face_aug_deg = read_json(METRICS / "face_resnet18_aug_test_degraded_metrics.json")
    full = read_json(METRICS / "full_image_resnet50_metrics.json")
    fusion = read_json(METRICS / "fusion_mlp_metrics.json")
    clip = read_json(METRICS / "clip_multimodal_metrics.json")
    return face, face_deg, face_aug, face_aug_deg, full, fusion, clip


def main():
    FIG_OUT.mkdir(parents=True, exist_ok=True)
    master = pd.read_csv(DATA / "msctd_master.csv")
    face_meta = pd.read_csv(DATA / "face_metadata.csv")
    fusion_features = pd.read_csv(DATA / "fusion_features.csv")
    test = fusion_features[fusion_features["split"] == "test"]
    face, face_deg, face_aug, face_aug_deg, full, fusion, clip = metric_rows()

    final_table = table_png(
        "table_final_comparison",
        ["Model", "Unit", "Acc.", "Macro-F1"],
        [
            ["Majority baseline", "image", "0.379", "0.183"],
            ["Face ResNet18", "crop", f"{face['accuracy']:.3f}", f"{face['macro_f1']:.3f}"],
            ["Face ResNet18, degraded", "crop", f"{face_deg['accuracy']:.3f}", f"{face_deg['macro_f1']:.3f}"],
            ["Aug. face, clean", "crop", f"{face_aug['accuracy']:.3f}", f"{face_aug['macro_f1']:.3f}"],
            ["Aug. face, degraded", "crop", f"{face_aug_deg['accuracy']:.3f}", f"{face_aug_deg['macro_f1']:.3f}"],
            ["Full-image ResNet50", "image", f"{full['accuracy']:.3f}", f"{full['macro_f1']:.3f}"],
            ["Fusion MLP", "image", f"{fusion['accuracy']:.3f}", f"{fusion['macro_f1']:.3f}"],
            ["CLIP image+text", "image+text", f"{clip['accuracy']:.3f}", f"{clip['macro_f1']:.3f}"],
        ],
        [2.45, 0.85, 0.75, 0.95],
    )
    robustness = pd.read_csv(METRICS / "robustness_results.csv")
    robust_table = table_png(
        "table_robustness",
        ["Experiment", "Acc.", "Macro-F1"],
        [[r["experiment"].replace("->", " to "), f"{r['accuracy']:.3f}", f"{r['macro_f1']:.3f}"] for _, r in robustness.iterrows()],
        [2.5, 0.8, 0.9],
    )
    perclass_table = table_png(
        "table_per_class",
        ["Model", "Neg. F1", "Neu. F1", "Pos. F1"],
        [
            ["Face ResNet18", f"{face['per_class']['negative']['f1']:.3f}", f"{face['per_class']['neutral']['f1']:.3f}", f"{face['per_class']['positive']['f1']:.3f}"],
            ["Full image", f"{full['per_class']['negative']['f1']:.3f}", f"{full['per_class']['neutral']['f1']:.3f}", f"{full['per_class']['positive']['f1']:.3f}"],
            ["Fusion MLP", f"{fusion['per_class']['negative']['f1']:.3f}", f"{fusion['per_class']['neutral']['f1']:.3f}", f"{fusion['per_class']['positive']['f1']:.3f}"],
            ["CLIP", f"{clip['per_class']['negative']['f1']:.3f}", f"{clip['per_class']['neutral']['f1']:.3f}", f"{clip['per_class']['positive']['f1']:.3f}"],
        ],
        [1.55, 0.8, 0.8, 0.8],
    )
    bar = comparison_bar([
        ("Majority baseline", 0.1833),
        ("Face ResNet18", face["macro_f1"]),
        ("Full-image ResNet50", full["macro_f1"]),
        ("Fusion MLP", fusion["macro_f1"]),
        ("CLIP image+text", clip["macro_f1"]),
    ])

    doc = Document()
    style_doc(doc)

    # Full-width title block.
    title = p(doc, "Multimodal Visual Sentiment Analysis on MSCTD: Faces, Scenes, Fusion, and Text", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_after = Pt(1)
    p(doc, "Shaik Sameer, Ammineni Harshitha, Tharigopula Sai Teja, Kakumanu Ravi Teja, Mosi Curran", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, "EEEM068 Applied Machine Learning, University of Surrey", size=8, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    abst = doc.add_paragraph()
    abst.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abst.paragraph_format.space_after = Pt(3)
    r = abst.add_run("Abstract—")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    r = abst.add_run(
        "This work studies three-class sentiment classification on the English-German split of MSCTD. "
        "We evaluate a staged pipeline comprising MTCNN face extraction, a ResNet18 face classifier, "
        "controlled robustness tests, a frozen ResNet50 full-image model, late fusion over visual probabilities, "
        "and an extra-credit CLIP image-text classifier. The processed corpus contains 30,370 image-text pairs and "
        f"{len(face_meta):,} detected face crops. Purely visual models occupy a narrow macro-F1 band near 0.36-0.38, "
        "showing that still-frame visual sentiment is intrinsically ambiguous in dialogue data. Augmentation reduces "
        "the face-model degradation gap from 5.5 to 0.5 macro-F1 points. Late fusion improves the negative class "
        "but not overall macro-F1, so it is best interpreted as a targeted routing result rather than a global win. "
        f"The CLIP multimodal branch is strongest, reaching {clip['accuracy']:.3f} accuracy and {clip['macro_f1']:.3f} macro-F1."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    p(doc, "Index Terms—sentiment analysis, MSCTD, MTCNN, ResNet, CLIP, multimodal fusion, robustness.", size=8, italic=True)

    # Two-column body.
    sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
    sec2.top_margin = Inches(0.62)
    sec2.bottom_margin = Inches(0.62)
    sec2.left_margin = Inches(0.62)
    sec2.right_margin = Inches(0.62)
    set_columns(sec2, 2, 360)

    heading(doc, "I. Introduction", 1)
    p(doc, "Visual sentiment analysis is difficult because a still frame often lacks the semantic context that makes an utterance negative, neutral, or positive. In MSCTD, dialogue frames can show neutral faces while the associated text expresses anger, irony, affection, or reassurance. This report therefore treats the coursework pipeline as a modality study: how far can face crops and full images go, and how much does text change the result?")
    p(doc, "The contribution is a reproducible seven-notebook pipeline. The study uses scene-level stratification to avoid dialogue leakage, explicit no-face and multi-face handling, controlled degradation tests, a frozen full-image transfer model, a compact late-fusion MLP, and a CLIP image-text model for the extra-credit component.")

    heading(doc, "II. Related Work", 1)
    p(doc, "ResNet remains a strong transfer-learning backbone because residual connections enable deeper convolutional networks to train reliably [1]. MTCNN is a standard cascade for face detection and alignment [2], while FaceNet illustrates the value of learned face representations [3]. Multimodal machine learning is commonly framed in terms of alignment and fusion [4]. CLIP is especially relevant because contrastive image-text pretraining yields aligned representations that can be reused with a lightweight classifier [5].")

    heading(doc, "III. Dataset and Preprocessing", 1)
    p(doc, f"The English-German split was parsed into {len(master):,} utterance-image pairs across 3,079 dialogue scenes. Splitting was performed at scene level, not row level, so related utterances from the same dialogue cannot appear in both training and test partitions. The final split contains 21,367 train, 4,572 validation, and 4,431 test samples. Neutral is the largest class at 39.0%, followed by negative at 32.9% and positive at 28.1%; consequently, macro-F1 is the primary metric.")
    add_fig(doc, PLOTS / "class_distribution.png", "Fig. 1. Class distribution in the processed MSCTD En-De split.")

    heading(doc, "IV. Methodology", 1)
    heading(doc, "A. Face Branch", 2)
    p(doc, f"MTCNN extracted {len(face_meta):,} face crops. In the test split, {int((test.num_faces > 0).sum()):,} of {len(test):,} images contain at least one detected face, while {int((test.num_faces == 0).sum()):,} contain none. Each crop inherits the image-level sentiment label. A ResNet18 initialized from ImageNet weights is fine-tuned on crop-level examples with inverse-frequency class weights. Multi-face predictions are averaged as softmax probabilities at image level.")
    add_fig(doc, PLOTS / "face_stats.png", "Fig. 2. Face-count and MTCNN confidence summaries.")

    heading(doc, "B. Robustness", 2)
    p(doc, "The robustness study applies three degradation families: brightness/contrast shifts, spatial transforms, and frequency corruptions such as blur, JPEG recompression, and Gaussian noise. Experiments A and B isolate the clean-to-degraded generalisation gap; experiment C retrains with mixed original/degraded inputs using p_degrade=0.5.")
    add_fig(doc, PLOTS / "augmentation_examples.png", "Fig. 3. Example brightness, spatial, and frequency degradations.", width=3.2)

    heading(doc, "C. Full-Image and Fusion Branches", 2)
    p(doc, "The full-image model uses a frozen ImageNet ResNet50 as a 2048-dimensional feature extractor followed by a 512-128-3 MLP head. It captures scene context, body posture, props, and lighting that are absent from face crops. The fusion model is an 8-dimensional late-fusion MLP over face probabilities, full-image probabilities, normalized face count, and a face-detected flag.")

    heading(doc, "D. CLIP Extra Credit", 2)
    p(doc, "The extra-credit model uses frozen CLIP ViT-B/32 image and text encoders. Their 512-dimensional projections are concatenated and passed into a small MLP classifier. Only 295,683 of 151.6M parameters are trainable. A low-learning-rate last-block unfreeze was implemented as a stage-two option, but the smoke test reduced validation macro-F1, so the frozen encoder configuration is reported.")

    heading(doc, "V. Experiments and Results", 1)
    add_table_fig(doc, final_table, "Table I. Test-set comparison. CLIP is added from the saved JSON metrics, correcting the stale CSV.", width=3.35)
    add_fig(doc, bar, "Fig. 4. Macro-F1 summary across the main evaluated models.")

    heading(doc, "A. Robustness", 2)
    add_table_fig(doc, robust_table, "Table II. Robustness results for original and augmentation-trained face models.", width=3.35)
    p(doc, "The original face model drops from 0.378 to 0.322 macro-F1 under degraded testing, a 5.5-point loss. The augmentation-trained model scores 0.366 on clean crops and 0.361 on degraded crops, reducing the robustness gap to 0.5 points. The small clean-performance cost is therefore exchanged for materially better deployment robustness.")

    heading(doc, "B. Per-Class Analysis", 2)
    add_table_fig(doc, perclass_table, "Table III. Per-class F1. Fusion improves negative F1 but weakens positive F1.")
    p(doc, "The full-image branch is strongest on neutral (0.492 F1), consistent with scene context helping visually flat dialogue frames. Fusion should be framed carefully: negative F1 improves to 0.430, above both face and full-image branches, but positive F1 falls to 0.192. Thus the fusion result is a targeted negative-class and regime-specific gain rather than an overall macro-F1 victory.")
    add_fig(doc, CMS / "fusion_mlp_cm.png", "Fig. 5. Fusion confusion matrix: negative recall improves, positive recall collapses.")

    heading(doc, "C. CLIP Multimodal Result", 2)
    p(doc, f"CLIP reaches {clip['accuracy']:.3f} accuracy, {clip['macro_f1']:.3f} macro-F1, and {clip['weighted_f1']:.3f} weighted-F1. Its per-class F1 values are {clip['per_class']['negative']['f1']:.3f} negative, {clip['per_class']['neutral']['f1']:.3f} neutral, and {clip['per_class']['positive']['f1']:.3f} positive. It is the only model for which every class exceeds 0.5 F1.")
    add_fig(doc, CMS / "clip_multimodal_cm.png", "Fig. 6. CLIP confusion matrix: substantially more balanced than fusion.")
    add_fig(doc, PLOTS / "clip_training_curve.png", "Fig. 7. CLIP frozen-encoder training curve.")

    heading(doc, "VI. Discussion", 1)
    p(doc, "The central empirical finding is a visual ceiling: face crops, full images, and probability-level fusion all improve over the majority baseline, yet remain close to 0.36-0.38 macro-F1. This supports the interpretation that MSCTD sentiment is often carried by utterance semantics rather than by still-frame visual evidence alone. The CLIP result is therefore not only a stronger score; it diagnoses where the missing information resides.")
    p(doc, "The fusion result remains valuable despite its lower macro-F1. It shows that face-count-aware routing contains signal, especially for the negative class, but that a shared MLP over probabilities is too blunt for positive sentiment and no-face cases. Better follow-up designs include calibrated logit fusion, class-conditional gates, or attention over the sequence of detected face embeddings.")

    heading(doc, "VII. Conclusion", 1)
    p(doc, "This project demonstrates a complete and reproducible multimodal sentiment pipeline for MSCTD En-De. The face branch establishes a meaningful but limited visual signal, augmentation improves robustness, the full-image model contributes scene context, and late fusion exposes useful negative-class routing without improving global macro-F1. The CLIP branch is the strongest model because it adds utterance text, reaching 0.562 macro-F1 with balanced per-class performance. The main conclusion is that dialogue sentiment cannot be reliably solved from still images alone; robust performance requires multimodal evidence.")

    heading(doc, "References", 1)
    refs = [
        "[1] K. He et al., \"Deep residual learning for image recognition,\" CVPR, 2016.",
        "[2] K. Zhang et al., \"Joint face detection and alignment using multitask cascaded convolutional networks,\" IEEE Signal Processing Letters, 2016.",
        "[3] F. Schroff et al., \"FaceNet: A unified embedding for face recognition and clustering,\" CVPR, 2015.",
        "[4] T. Baltrusaitis et al., \"Multimodal machine learning: A survey and taxonomy,\" IEEE TPAMI, 2018.",
        "[5] A. Radford et al., \"Learning transferable visual models from natural language supervision,\" ICML, 2021.",
        "[6] Y. Liang et al., \"MSCTD: A multimodal sentiment chat translation dataset,\" ACL, 2022.",
        "[7] A. Dosovitskiy et al., \"An image is worth 16x16 words,\" ICLR, 2021.",
        "[8] Z. Liu et al., \"Swin Transformer,\" ICCV, 2021.",
        "[9] H. Tan and M. Bansal, \"LXMERT,\" EMNLP, 2019.",
        "[10] J. Lei et al., \"Less is more: ClipBERT,\" CVPR, 2021.",
    ]
    for ref in refs:
        p(doc, ref, size=7.5)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
